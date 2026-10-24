"""
Tests for clustered candidate files (Sector Surveyor 1.6.0+): format
detection, survey-region legs, claim filter, resolution, live-edit merge,
next target, region exports, and the migration commands.
"""
from __future__ import annotations

import csv
import io
import json
import os
from datetime import datetime, timezone

import openpyxl
import pytest

from engine.db       import Database
from engine.exporter import export_for_config
from engine.legs     import LegManager
from engine.parser   import EventParser
from engine.survey   import SurveyFormatError, SurveyTracker, detect_format
from tests.test_survey import all_found, jump, read_rows, star_scan

CLUSTERED_HEADER = [
    "visit_order", "survey_region", "region_visit", "cluster", "cluster_stop", "boxel_id",
    "system_name", "candidate_type", "edsm_status", "direction", "family", "steps_from_edge",
    "subsector", "mass_code", "mass_prefix", "number", "boxel_known_count", "boxel_x",
    "boxel_y", "boxel_z", "cube_min_x", "cube_min_y", "cube_min_z", "cube_size_ly",
    "anchor_known_system",
]
A1, A2 = "Hypiae Aim AA-A c1-1", "Hypiae Aim AA-A c1-2"   # R01
B1, B2 = "Hypiae Aim BB-B c1-1", "Hypiae Aim BB-B c1-2"   # R02
LOST   = "Hypiae Aim CC-C c1-1"                           # no survey region
OUTSIDE = "Hypiae Aim ZZ-Z c1-9"                          # not a candidate


def _row(visit_order, region, name):
    p = name.split()
    return {
        "visit_order": visit_order, "survey_region": region,
        "region_visit": "", "cluster": f"{region}-C01" if region else "",
        "cluster_stop": "1", "boxel_id": f"{p[2]} c1", "system_name": name,
        "candidate_type": "bracketed_gap", "edsm_status": "not_in_edsm",
        "direction": "bracketed_gap", "family": f"Hypiae Aim {p[2]}",
        "steps_from_edge": "", "subsector": p[2], "mass_code": "c", "mass_prefix": "c1",
        "number": p[3].split("-")[1], "boxel_known_count": "3", "boxel_x": "-8000.0",
        "boxel_y": "100.0", "boxel_z": "-5600.0", "cube_min_x": "-8005.0",
        "cube_min_y": "95.0", "cube_min_z": "-5605.0", "cube_size_ly": "40",
        "anchor_known_system": "Hypiae Aim AA-A c1-0",
    }


def write_clustered(path, rows=None):
    rows = rows or [_row("1", "R01", A1), _row("2", "R01", A2),
                    _row("3", "R02", B1), _row("4", "R02", B2), _row("", "", LOST)]
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=CLUSTERED_HEADER, lineterminator="\r\n")
    w.writeheader()
    w.writerows(rows)
    path.write_bytes(buf.getvalue().encode("utf-8"))   # no BOM, CRLF — as from Excel
    return path


def bump_mtime(path):
    st = path.stat()
    os.utime(path, ns=(st.st_atime_ns, st.st_mtime_ns + 5_000_000_000))


@pytest.fixture
def csv_path(tmp_path):
    return write_clustered(tmp_path / "hypiae_aim_candidates_clustered.csv")


@pytest.fixture
def sector_cfg(cfg, tmp_path, csv_path):
    c = dict(cfg)
    c.update(commanders=["AristonX", "DovakiinMJ"], mode="sector", sector="Hypiae Aim",
             waypoints=[], waypoint_map={}, expedition_end_system="",
             expedition_end_system_norm="", output_dir=tmp_path / "out", rarity={},
             db_path=tmp_path / "test.db",
             survey={"csv_path": csv_path, "region_complete_when": "completed"})
    c["output_dir"].mkdir()
    return c


class Rig:
    """Parser + legs + survey on one DB, with a journal clock."""

    def __init__(self, db, cfg, survey):
        self.db, self.survey = db, survey
        self.legs = LegManager(db, cfg)
        self.legs.ensure_first_leg()
        self.closed: list[int] = []
        self.parser = EventParser(db, self.legs, cfg, survey=survey,
                                  on_leg_close=self.closed.append)
        self.parser.restore_state()
        self._t = 0

    def ts(self):
        self._t += 1
        return f"2026-10-25T{self._t // 60:02d}:{self._t % 60:02d}:00Z"

    def feed(self, ev):
        ev = dict(ev, timestamp=self.ts())
        self.parser.process_line(json.dumps(ev))

    def login(self, name="AristonX"):
        self.feed({"event": "Commander", "FID": "F1", "Name": name})

    def jump(self, name):
        self.feed(jump(name))

    def explore(self, name, discovered=False):
        self.jump(name)
        self.feed(star_scan(name, discovered))
        self.feed(all_found(name))

    def leg_names(self):
        return [leg["name"] for leg in self.db.get_all_legs()]


def make_rig(db, cfg, **survey_kw):
    s = SurveyTracker(cfg["survey"]["csv_path"], "Hypiae Aim", retry_seconds=3600, **survey_kw)
    rig = Rig(db, cfg, s)
    rig.login()
    return rig


# ── Format detection ───────────────────────────────────────────────────────────

def test_detects_clustered_format_and_keeps_file_style(csv_path):
    s = SurveyTracker(csv_path, "Hypiae Aim")
    assert s.format == "clustered" and s.clustered
    assert s.regions == ["R01", "R02"]
    s.on_event(jump(A1), "AristonX")
    raw = csv_path.read_bytes()
    assert not raw.startswith(b"\xef\xbb\xbf")
    assert b"\r\n" in raw
    header = raw.decode("utf-8").splitlines()[0].split(",")
    assert header == CLUSTERED_HEADER + [
        "discovered", "completed", "region", "visited_utc", "completed_utc", "resolution"]
    assert [r["system_name"] for r in read_rows(csv_path).values()] == [A1, A2, B1, B2, LOST]


def test_detect_format_plain():
    fields = ["system_name", "edsm_status", "direction", "family"]
    assert detect_format(fields, [{"system_name": "x"}]) == "plain"
    # A plain file the tracker already wrote has a galactic `region` column.
    assert detect_format(fields + ["region"], [{"region": "Elysian Shore"}]) == "plain"


def test_pre_160_region_column_is_refused(tmp_path):
    p = tmp_path / "old_clustered.csv"
    p.write_text("visit_order,region,cluster,system_name\n1,R01,R01-C01,Hypiae Aim AA-A c1-1\n",
                 encoding="utf-8")
    before = p.read_bytes()
    with pytest.raises(SurveyFormatError, match="1.6.0"):
        SurveyTracker(p, "Hypiae Aim")
    assert p.read_bytes() == before


# ── Region legs ────────────────────────────────────────────────────────────────

def test_region_closes_on_next_jump_away(db, sector_cfg):
    rig = make_rig(db, sector_cfg)
    rig.explore(A1)
    rig.explore(A2)                      # R01 now complete, still in A2
    assert rig.closed == []
    assert db.get_active_leg()["end_ts"] is None

    rig.jump(OUTSIDE)                    # depart A2 → close
    assert len(rig.closed) == 1
    closed = db.get_leg(rig.closed[0])
    assert closed["name"] == "R01"
    assert closed["end_system"] == A2
    assert db.get_active_leg() is not None
    assert "R01" in db.get_visited_waypoints()
    # The departing jump belongs to the new leg.
    assert db.get_last_jump(db.get_active_leg()["leg_id"])["to_system"] == OUTSIDE


def test_visited_mode_does_not_close_on_arrival(db, sector_cfg):
    rig = make_rig(db, sector_cfg, complete_when="visited")
    rig.jump(A1)
    rig.jump(A2)                         # arrival completes R01 — no close yet
    assert rig.closed == []
    rig.jump(OUTSIDE)
    assert rig.leg_names()[0] == "R01"


def test_regions_any_order_then_expedition_complete(db, sector_cfg):
    rig = make_rig(db, sector_cfg)
    rig.explore(B1)
    rig.explore(B2)
    rig.explore(A1)                      # departs B2 → R02 closes
    assert rig.leg_names() == ["R02", None]
    rig.explore(A2)
    rig.jump(OUTSIDE)                    # the no-region row doesn't block completion
    assert rig.leg_names() == ["R02", "R01"]
    assert rig.legs.is_complete
    assert db.get_active_leg() is None
    rig.explore(LOST)                    # still tracked after completion
    assert read_rows(sector_cfg["survey"]["csv_path"])[LOST]["completed"] == "AristonX"


def test_claim_filter(db, sector_cfg):
    rig = make_rig(db, sector_cfg, regions=["r02"])
    assert rig.survey.regions == ["R02"]
    rig.explore(A1)
    rig.explore(A2)
    rig.jump(OUTSIDE)                    # R01 isn't claimed: no leg
    assert rig.closed == []
    rig.explore(B1)
    rig.explore(B2)
    rig.jump(OUTSIDE)
    assert rig.leg_names() == ["R02"]
    assert rig.legs.is_complete
    rows = read_rows(sector_cfg["survey"]["csv_path"])
    assert rows[A1]["completed"] == "AristonX"     # unclaimed rows still tracked


def test_two_regions_finishing_together_close_one_leg(db, sector_cfg):
    rig = make_rig(db, sector_cfg)
    rig.survey.resolve(A1, "not_found")
    rig.survey.resolve(A2, "not_found")
    rig.survey.resolve(B1, "done_by_other")
    rig.survey.resolve(B2, "not_found")
    rig.survey.flush()
    rig.jump(OUTSIDE)                    # no current system yet: nothing to depart
    rig.jump(LOST)
    assert rig.leg_names() == ["R01 + R02"]
    assert db.get_visited_waypoints() >= {"R01", "R02"}
    assert rig.legs.is_complete


def test_not_found_resolution_closes_region(db, sector_cfg):
    rig = make_rig(db, sector_cfg)
    rig.explore(A1)
    rig.survey.resolve(A2, "not_found")
    rig.survey.flush()
    rig.jump(OUTSIDE)
    assert rig.leg_names()[0] == "R01"
    s = rig.survey.summary()
    assert s["not_found"] == 1 and s["resolved"] == 2 and s["regions_complete"] == 1


def test_pending_region_closes_after_restart_and_only_once(tmp_path, sector_cfg):
    db1 = Database(sector_cfg["db_path"])
    rig = make_rig(db1, sector_cfg)
    rig.explore(A1)
    rig.explore(A2)
    rig.survey.close()
    db1.close()

    db2 = Database(sector_cfg["db_path"])
    rig2 = Rig(db2, sector_cfg, SurveyTracker(sector_cfg["survey"]["csv_path"], "Hypiae Aim"))
    rig2._t = 100
    rig2.jump(OUTSIDE)
    rig2.jump(B1)
    assert rig2.leg_names() == ["R01", None]
    db2.close()


def test_plain_file_never_closes_region_legs(db, sector_cfg, tmp_path):
    plain = tmp_path / "plain.csv"
    plain.write_text("system_name,edsm_status,direction\n"
                     f"{A1},not_in_edsm,backward\n", encoding="utf-8")
    s = SurveyTracker(plain, "Hypiae Aim")
    assert s.format == "plain" and s.regions == [] and s.next_target() is None
    rig = Rig(db, sector_cfg, s)
    rig.explore(A1)
    rig.jump(OUTSIDE)
    assert rig.closed == [] and rig.leg_names() == ["Hypiae Aim"]
    assert "resolution" not in plain.read_text(encoding="utf-8").splitlines()[0]
    with pytest.raises(SurveyFormatError):
        s.resolve(A1, "not_found")


# ── Live-edit merge ────────────────────────────────────────────────────────────

def _edit(path, edits, add_col=None):
    """Rewrite the CSV as an external editor would: {system: {col: value}}."""
    text = path.read_bytes().decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text, newline=""))
    fields = list(reader.fieldnames) + ([add_col] if add_col else [])
    rows = list(reader)
    for cols in edits.values():
        fields += [c for c in cols if c not in fields]
    for r in rows:
        r.update(edits.get(r["system_name"], {}))
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=fields, lineterminator="\r\n")
    w.writeheader()
    w.writerows(rows)
    path.write_bytes(buf.getvalue().encode("utf-8"))
    bump_mtime(path)


def test_live_edit_merge(csv_path):
    s = SurveyTracker(csv_path, "Hypiae Aim", retry_seconds=3600)
    s.on_event(jump(A1), "AristonX")
    s.on_event(all_found(A1), "AristonX")

    # Commander marks A2 not found; Excel mangles A1's timestamps and they add a column.
    _edit(csv_path, {A2: {"resolution": "not_found", "notes": "checked galmap"},
                     A1: {"visited_utc": "10/23/2026 1:00", "completed": ""}},
          add_col="notes")
    s.on_event(jump(B1), "AristonX")     # triggers a write

    rows = read_rows(csv_path)
    assert rows[A2]["resolution"] == "not_found"      # disk wins for resolution
    assert rows[A2]["notes"] == "checked galmap"      # user's column kept
    assert rows[A1]["visited_utc"] == "2026-10-23T01:00:00Z"   # tracker wins for its own
    assert rows[A1]["completed"] == "AristonX"
    assert rows[B1]["visited_utc"]                    # the new event landed too
    assert s.complete_regions() == ["R01"]

    # Clearing a mistaken resolution on disk is honoured.
    _edit(csv_path, {A2: {"resolution": ""}})
    assert s.complete_regions() == []
    s.close()
    assert read_rows(csv_path)[A2]["resolution"] == ""


def test_edit_fills_empty_tracker_cell(csv_path):
    s = SurveyTracker(csv_path, "Hypiae Aim")
    _edit(csv_path, {B2: {"completed": "Squadmate"}})
    assert s.complete_regions() == []
    s.on_event(jump(B1), "AristonX")
    s.on_event(all_found(B1), "AristonX")
    assert s.complete_regions() == ["R02"]
    assert read_rows(csv_path)[B2]["completed"] == "Squadmate"


# ── Serendipitous rows ─────────────────────────────────────────────────────────

def test_serendipitous_on_csv_without_edsm_status_or_direction(tmp_path):
    p = tmp_path / "minimal.csv"
    p.write_text("system_name,family,subsector\n" f"{A1},Hypiae Aim AA-A,AA-A\n", encoding="utf-8")
    s = SurveyTracker(p, "Hypiae Aim")
    s.on_event(jump("Hypiae Aim ZI-B d4"), "AristonX")
    s.on_event(star_scan("Hypiae Aim ZI-B d4", False), "AristonX")   # flushes — must not raise
    row = read_rows(p)["Hypiae Aim ZI-B d4"]
    assert row["family"] == "Hypiae Aim ZI-B" and row["discovered"] == "new"
    assert "direction" not in row
    assert s.summary()["visited"] == 1                                 # no KeyError
    assert "Survey:" in s.summary_text()


def test_serendipitous_in_clustered_file_has_no_survey_region(csv_path):
    s = SurveyTracker(csv_path, "Hypiae Aim")
    s.on_event(jump(OUTSIDE), "AristonX")
    s.on_event(star_scan(OUTSIDE, False), "AristonX")
    rows = read_rows(csv_path)
    row = rows[OUTSIDE]
    assert list(rows)[-1] == OUTSIDE
    assert (row["survey_region"], row["visit_order"], row["direction"], row["boxel_id"]) == \
        ("", "", "serendipitous", "ZZ-Z c1")
    assert s.region_progress() == {"R01": (0, 2), "R02": (0, 2)}


# ── Progress / next target ─────────────────────────────────────────────────────

def test_next_target_follows_active_region(csv_path):
    s = SurveyTracker(csv_path, "Hypiae Aim")
    assert s.next_target()["visit_order"] == "1"
    s.on_event(jump(B1), "AristonX")
    s.on_event(all_found(B1), "AristonX")
    assert s.next_target()["system_name"] == B2       # active region R02 first
    text = s.summary_text()
    assert "R02: 1/2 resolved" in text and f"next #4 {B2}" in text
    assert s.short_status() == "R02 1/2"

    s.on_event(jump(B2), "AristonX")
    s.on_event(all_found(B2), "AristonX")
    assert s.next_target()["system_name"] == A1       # falls back to other claimed region
    assert s.short_status() == "R01 0/2"
    # Active region is restored from the CSV after a restart.
    s.close()
    assert SurveyTracker(csv_path, "Hypiae Aim")._active_region == "R02"


# ── Exports ────────────────────────────────────────────────────────────────────

def test_region_leg_export_has_candidates_sheet(db, sector_cfg):
    rig = make_rig(db, sector_cfg)
    rig.explore(A1)
    rig.explore(A2)
    rig.jump(OUTSIDE)
    closed = db.get_leg(rig.closed[0])
    xlsx = export_for_config(sector_cfg, db, closed["leg_id"], survey=rig.survey)
    assert xlsx.name == f"sector_hypiae_aim_r01_{closed['end_ts'][:10]}.xlsx"
    wb = openpyxl.load_workbook(xlsx)
    assert "Candidates" in wb.sheetnames
    cand = (sector_cfg["output_dir"] / f"{xlsx.stem}_candidates.csv").read_text().splitlines()
    assert len(cand) == 3 and A1 in cand[1] and A2 in cand[2]

    # The open leg keeps the cumulative as-of name and no Candidates sheet.
    open_x = export_for_config(sector_cfg, db, db.get_active_leg()["leg_id"], survey=rig.survey)
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    assert open_x.name == f"sector_hypiae_aim_{today}.xlsx"
    assert "Candidates" not in openpyxl.load_workbook(open_x).sheetnames


# ── Migration commands ─────────────────────────────────────────────────────────

def test_migration_close_leg_then_rebuild(sector_cfg, tmp_path, capsys):
    """Plain-list survey in progress → close its leg by hand → switch csv_path
    to the clustered file → --survey-rebuild → region legs from then on."""
    import tracker

    plain = tmp_path / "plain.csv"
    plain.write_text("system_name,edsm_status,direction\n"
                     f"{A1},not_in_edsm,backward\n", encoding="utf-8")
    plain_cfg = dict(sector_cfg, survey={"csv_path": plain})
    db = Database(plain_cfg["db_path"])
    rig = make_rig(db, plain_cfg)
    rig.explore(A1)
    rig.survey.close()
    db.close()

    tracker._run_close_leg(plain_cfg, "Hypiae Aim (pre-regions)")
    out = capsys.readouterr().out
    assert "closed as 'Hypiae Aim (pre-regions)'" in out
    db = Database(plain_cfg["db_path"])
    first = db.get_all_legs()[0]
    assert first["name"] == "Hypiae Aim (pre-regions)" and first["end_ts"]
    assert db.get_active_leg()["ordinal"] == 2
    db.close()
    date = first["end_ts"][:10]
    assert (sector_cfg["output_dir"] / f"sector_hypiae_aim_hypiae_aim_pre_regions_{date}.xlsx").exists()

    # Point the config at the clustered file and replay progress into it.
    tracker._run_survey_rebuild(sector_cfg)
    rows = read_rows(sector_cfg["survey"]["csv_path"])
    assert rows[A1]["completed"] == "AristonX"

    # Finishing R01 now closes the second leg as "R01", not the pre-regions one.
    db = Database(sector_cfg["db_path"])
    rig2 = Rig(db, sector_cfg, SurveyTracker(sector_cfg["survey"]["csv_path"], "Hypiae Aim"))
    rig2._t = 500
    rig2.explore(A2)
    rig2.jump(OUTSIDE)
    assert rig2.leg_names() == ["Hypiae Aim (pre-regions)", "R01", None]
    db.close()


def test_survey_resolve_command(sector_cfg, capsys):
    import tracker
    tracker._run_survey_resolve(sector_cfg, [A1, "hypiae aim aa-a c1-2"], "not_found")
    rows = read_rows(sector_cfg["survey"]["csv_path"])
    assert rows[A1]["resolution"] == rows[A2]["resolution"] == "not_found"
    assert "2 system(s) set to 'not_found'" in capsys.readouterr().out

    tracker._run_survey_resolve(sector_cfg, [A1], "clear")
    assert read_rows(sector_cfg["survey"]["csv_path"])[A1]["resolution"] == ""

    with pytest.raises(SystemExit):
        tracker._run_survey_resolve(sector_cfg, ["Hypiae Aim QQ-Q c9-9"], "not_found")
